# -*- coding: utf-8 -*-
# 2026-10-02 작업 결과를 워드 문서로 만든다.
# 출력: E:/AI 프로젝트 결과물/project table 정리/홈페이지_프로젝트_오늘의_진행상황_Vercel이전_애드센스_2026-10-02.docx
# 주의: 토큰·비밀번호 등 민감정보는 절대 넣지 않는다 (README 운영 원칙).

import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = r"E:\AI 프로젝트 결과물\project table 정리"
OUT = os.path.join(OUT_DIR, "홈페이지_프로젝트_오늘의_진행상황_Vercel이전_애드센스_2026-10-02.docx")

doc = Document()

# 기본 폰트를 맑은 고딕으로
style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10.5)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    return par


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = hd
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    doc.add_paragraph()


# ── 표지 ─────────────────────────────────────────────
title = doc.add_heading("홈페이지·프로젝트 오늘의 진행상황", level=0)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Vercel 프로젝트 이전 · 배포 장애 복구 · 애드센스 상태 (2026-10-02)")
r.bold = True
r.font.size = Pt(12)
p("대상: 한국노총 노동연대 홈페이지 (fktu-nodong) · 하루기록 블로그 (harugirok)")
p("※ 이 문서에는 토큰·비밀번호 등 민감정보를 담지 않았습니다. 접속 정보는 별도 문서에서 관리합니다.")

# ── 1. 요약 ──────────────────────────────────────────
h("1. 오늘 한 것 요약", 1)
table(
    ["구분", "내용", "결과"],
    [
        ["배포 장애 진단", "구 프로젝트의 모든 배포가 빌드 시작 전 실패", "프로젝트 단위 고착 확인"],
        ["Vercel 프로젝트 이전", "fktu-nodong-v2 신규 프로젝트로 이전", "배포 READY · 주소 불변"],
        ["라이브 검증", "9개 페이지 · 링크 2,197개 전수 검사", "깨짐 0"],
        ["사이트 사본 재생성", "_사이트사본-2026-09-16/ 갱신", "110파일 · 원본 대조 차이 0"],
        ["GitHub 백업", "커밋 2건 푸시 · push hang 원인 해결", "main 동기화 완료"],
        ["Vercel 지원 티켓", "영문 초안 작성 (원인 확인용)", "보관 완료"],
        ["애드센스 준비 점검", "노동연대 사이트 요건 점검", "기본 요건 충족"],
        ["구 프로젝트 폐기", "고장 난 프로젝트 삭제", "정리 완료"],
    ],
)

# ── 2. Vercel 배포 장애 진단과 복구 ───────────────────
h("2. Vercel 배포 장애 진단과 복구", 1)
h("2.1 증상", 2)
p("구 프로젝트(fktu-nodong)로의 배포가 전부 아래와 같이 실패했습니다.")
table(
    ["시도", "방식", "결과"],
    [
        ["05:51", "GitHub 자동배포 (커밋 연동)", "BUILD_FAILED / Resource provisioning failed"],
        ["05:56", "Vercel CLI", "같은 오류"],
        ["06:04", "GitHub 자동배포", "같은 오류"],
        ["06:09", "API 직접 업로드", "같은 오류"],
    ],
)
p("빌드 시간 0ms, 빌드 로그 없음 → 빌드 컨테이너가 시작조차 되지 않은 상태로, 사이트 파일과는 무관한 장애입니다.")

h("2.2 진단", 2)
bullet("계정 상태: 정상 (active, hobby 플랜)")
bullet("공개 API로 바꿀 수 있는 설정을 모두 정리 (fluid 해제, 함수 타임아웃 300→10초, 배포 보호 해제) → 실패 동일")
bullet("같은 설정의 임시 신규 프로젝트는 즉시 READY 성공 → 계정·코드 정상, 구 프로젝트에만 고착된 비공개 레거시 설정이 원인")
bullet("동일 사례: Vercel 커뮤니티의 \"stuck legacy backend settings\" 사례 (공개 API로는 초기화 불가, 지원팀 조치 필요)")

h("2.3 조치 — fktu-nodong-v2 프로젝트로 이전", 2)
table(
    ["항목", "내용"],
    [
        ["신규 프로젝트", "fktu-nodong-v2 (프로젝트 ID prj_WUsBK4YXEtyS8tyjYnlBkI3lwfGE)"],
        ["배포", "113개 파일 업로드 → READY (배포 ID dpl_FHMMjrZTzVeBHhgqKq6uMy1oHwNA)"],
        ["도메인", "fktu-nodong.vercel.app 그대로 이전 (verified) — 주소 불변"],
        ["영향", "방문자 주소·기능 손실 없음. 애드센스 신청 주소도 그대로 유효"],
        ["구 프로젝트", "도메인 분리 확인 후 폐기 완료 (원인 기록은 지원 티켓 초안에 보관)"],
    ],
)

# ── 3. 라이브 검증 ───────────────────────────────────
h("3. 라이브 검증 결과", 1)
bullet("9개 페이지 전부 200 응답 (홈·소개·문의·환영·지식·소식·게시판·가입·개인정보)")
bullet("홈·소개에 다듬은 새 문구 반영 확인")
bullet("\"64편\" 표기 반영 확인 (지식 글 64편 기준)")
bullet("박강원 표기 8곳 유지 (공개 7 + README 1) — 기존 결정대로 유지하고 보고만 함")
bullet("내부 링크 전수 검사: HTML 86개 · 링크 2,197개 · 깨짐 0")
bullet("문서 파일(docx 등)이 라이브에 노출되지 않는지 확인 → 전부 404 (안전)")

# ── 4. 사이트 사본 ──────────────────────────────────
h("4. 사이트 사본 재생성", 1)
p("_사이트사본-2026-09-16/ 폴더를 최신 배포본으로 다시 만들었습니다 (이름은 유지).")
bullet("임시 폴더에 먼저 생성한 뒤 교체하는 방식으로 안전하게 처리")
bullet("결과: 110파일 · 원본과 대조하여 차이 0")

# ── 5. GitHub 백업 ──────────────────────────────────
h("5. GitHub 백업과 push hang 해결", 1)
bullet("커밋 2건 푸시 완료 (프로젝트 이전 기록 · 구 프로젝트 폐기 기록)")
bullet("문제: git push 가 10분 넘게 멈춤 (로컬 커밋은 정상)")
bullet("원인: 이 PC 의 기본 자격증명 헬퍼(manager, GCM) 가 대화상자에서 멈추는 문제")
bullet("해결: gh 토큰을 직접 쓰는 명령으로 즉시 푸시 성공 → README에 해결책 기록")
p("※ git ls-remote·fetch 는 정상이어서 네트워크 문제는 아닙니다.")

# ── 6. 지원 티켓 ───────────────────────────────────
h("6. Vercel 지원 티켓 초안", 1)
bullet("위치: fktu-nodong/tools/vercel-support-ticket.md (git 미추적 = 안전)")
bullet("목적: 복구 요청이 아닌 원인 확인·재발 방지 문의 (구 프로젝트는 폐기했기 때문)")
bullet("실패 배포 ID 4건, 시도한 설정 정리, 동일 사례 링크 포함 — 그대로 붙여넣어 접수 가능")

# ── 7. 애드센스 준비 점검 ───────────────────────────
h("7. 애드센스 승인 준비 상태 (노동연대 사이트)", 1)
h("7.1 준비 완료", 2)
table(
    ["항목", "상태"],
    [
        ["개인정보처리방침", "privacy.html 8개 섹션 완비"],
        ["이용약관·연락처·소개", "terms.html · contact.html · about.html"],
        ["푸터 접근성", "소개·문의·개인정보처리방침·이용약관 링크 완비"],
        ["콘텐츠", "지식 64편 + 소식, 전건 출처 표기 (저작물 문제 없음)"],
        ["기본 인프라", "HTTPS · robots.txt · sitemap.xml · Google 소유 확인 파일"],
        ["민감 파일", "라이브 미노출(404) 확인"],
    ],
)
h("7.2 남은 것", 2)
table(
    ["항목", "내용"],
    [
        ["ads.txt", "없음 — 애드센스 승인 시 발급되는 pub ID로 생성해야 함"],
        ["개인정보처리방침 7절", "현재 \"광고 쿠키를 설치하지 않습니다\"인데 애드센스 적용 시 사실과 달라짐 → Google 광고 쿠키 고지로 개정 필요"],
        ["광고 코드", "승인 후 <head> 에 애드센스 스니펫 삽입"],
    ],
)

# ── 8. 애드센스 심사 결과 ───────────────────────────
h("8. 애드센스 심사 결과 (하루기록 블로그)", 1)
p("심사 결과: 거절 — 사유 \"가치가 별로 없는 콘텐츠(Content with little value)\"", bold=True)
h("8.1 이미 조치한 것 (배포 완료)", 2)
bullet("소개 페이지에 \"글은 이렇게 만듭니다\" (현장 경험 → 출처 확인 → 사람 검토 → 수정일 표시)")
bullet("정정과 업데이트 원칙, 광고와 편집의 분리 선언 추가")
bullet("글 페이지에 글쓴이 표시 + \"사람 검토 후 게시\" 바이라인 (자동 생성 글이 아님을 보여주는 신호)")
h("8.2 권장 순서", 2)
p("거절 사유의 본질은 글의 양·깊이·트래픽이라, 바로 재심사 요청 시 재거절 가능성이 높습니다.")
bullet("1) \"깊은 글\" 추가 — 계산기·체크리스트·비교표처럼 체류 시간이 긴 실용 콘텐츠")
bullet("2) 검색 유입 확보 — 색인이 시작된 단계라 2~4주 자연 유입 지켜보기")
bullet("3) 그다음 \"검토 요청\" 버튼 클릭")

# ── 9. 최종 주소 ───────────────────────────────────
h("9. 최종 주소", 1)
table(
    ["사이트", "주소", "상태"],
    [
        ["한국노총 노동연대", "https://fktu-nodong.vercel.app", "정상 (주소 불변)"],
        ["하루기록 블로그", "https://harugirok.vercel.app", "애드센스 거절 상태"],
    ],
)

# ── 10. 남은 할 일 ──────────────────────────────────
h("10. 남은 할 일", 1)
table(
    ["할 일", "대상", "비고"],
    [
        ["애드센스 재심사 대응 (깊은 글 작성)", "하루기록", "거절 사유 본질 대응"],
        ["ads.txt 준비 + 개인정보처리방침 개정", "노동연대", "승인 시점에 맞춰"],
        ["Vercel 지원 티켓 접수", "노동연대", "초안 있음 — 원인 확인용"],
        ["새 프로젝트 GitHub 연결 (자동배포)", "노동연대", "필요 시 대시보드에서"],
    ],
)

doc.save(OUT)
print("저장 완료:", OUT)
print("크기:", os.path.getsize(OUT), "bytes")
